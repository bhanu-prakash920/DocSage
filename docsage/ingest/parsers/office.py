"""Word (.docx) and PowerPoint (.pptx) parsers."""

from __future__ import annotations

from pathlib import Path

from ..schema import ImageElement, PageText, ParsedDocument, TableElement
from .base import CAPTION_RE, ParseContext, rows_to_markdown, tidy
from .sections import split_sections

_BLIP = ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def parse_docx(path: Path, ctx: ParseContext) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    title = (document.core_properties.title or path.stem).strip() or path.stem
    result = ParsedDocument(title=title, kind="docx", unit="section")

    lines: list[str] = []
    pending_tables: list[tuple[int, str]] = []  # (char offset, markdown)
    pending_images: list[tuple[int, bytes, str]] = []
    last_text = ""
    awaiting_caption: int | None = None  # index into pending_images of the last image seen
    captions: dict[int, str] = {}
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        offset = sum(len(x) + 1 for x in lines)
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            style = (para.style.name if para.style is not None else "") or ""
            for blip in child.findall(_BLIP):
                rid = blip.get(_EMBED)
                part = document.part.related_parts.get(rid) if rid else None
                if part is not None and getattr(part, "blob", None):
                    pending_images.append((offset, part.blob, last_text))
                    awaiting_caption = len(pending_images) - 1
            if not text:
                continue
            if awaiting_caption is not None:
                # Captions usually follow the figure; accept the first paragraph after it if it looks like one.
                if CAPTION_RE.match(text):
                    captions[awaiting_caption] = text
                awaiting_caption = None
            if style.startswith("Heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "2"
                lines.append("#" * min(int(level), 4) + " " + text)
            elif style.startswith("List"):
                lines.append(f"- {text}")
            else:
                lines.append(text)
            last_text = text
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            md = rows_to_markdown(rows)
            if md:
                lines.append(md)
                pending_tables.append((offset, md))

    markdown = tidy("\n\n".join(lines))
    result.pages = split_sections(markdown)
    from .sections import section_of

    for offset, md in pending_tables:
        result.tables.append(TableElement(page=section_of(offset, result.pages), markdown=md))
    for index, (offset, blob, before) in enumerate(pending_images):
        saved = ctx.save_png(blob, f"s{section_of(offset, result.pages)}_img{index + 1}")
        if saved:
            caption = captions.get(index) or (before if CAPTION_RE.match(before or "") else "")
            result.images.append(
                ImageElement(
                    page=section_of(offset, result.pages),
                    image_path=saved[0],
                    caption=caption,
                    context=before[:400],
                    width=saved[1],
                    height=saved[2],
                )
            )
    ctx.report(1.0, "Parsed")
    return result


def parse_pptx(path: Path, ctx: ParseContext) -> ParsedDocument:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(str(path))
    title = (deck.core_properties.title or path.stem).strip() or path.stem
    result = ParsedDocument(title=title, kind="pptx", unit="slide")
    total = len(deck.slides)
    for index, slide in enumerate(deck.slides, start=1):
        ctx.report((index - 1) / max(1, total), f"Parsing slide {index} of {total}")
        parts: list[str] = []
        slide_title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            slide_title = slide.shapes.title.text_frame.text.strip()
            if slide_title:
                parts.append(f"# {slide_title}")
        for s_index, shape in enumerate(slide.shapes):
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if text.strip():
                    parts.append(text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                md = rows_to_markdown(rows)
                if md:
                    parts.append(md)
                    result.tables.append(TableElement(page=index, markdown=md, caption=slide_title))
            if getattr(shape, "has_chart", False) and shape.has_chart:
                md = _chart_markdown(shape.chart)
                if md:
                    result.tables.append(
                        TableElement(
                            page=index, markdown=md, caption=f"Chart data: {slide_title}".strip(": ")
                        )
                    )
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                saved = ctx.save_png(shape.image.blob, f"s{index}_img{s_index + 1}")
                if saved:
                    result.images.append(
                        ImageElement(
                            page=index,
                            image_path=saved[0],
                            caption=slide_title,
                            context=" ".join(parts)[:400],
                            width=saved[1],
                            height=saved[2],
                        )
                    )
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        result.pages.append(PageText(index, tidy("\n\n".join(parts))))
    ctx.report(1.0, "Parsed")
    return result


def _chart_markdown(chart) -> str:
    try:
        plot = chart.plots[0]
        categories = [str(c) for c in plot.categories]
        rows = [["Series", *categories]]
        for series in plot.series:
            rows.append([series.name or "", *[("" if v is None else f"{v:g}") for v in series.values]])
        return rows_to_markdown(rows)
    except Exception:
        return ""
