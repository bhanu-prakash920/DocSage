"""Generates a small, fictional, multimodal sample corpus and its golden question set.

Used by the "Load sample collection" onboarding action, the test suite and the benchmark.
All companies, people and figures are invented.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFont

REPORT = "halcyon-grid-annual-report-2025.pdf"
FIELD_GUIDE = "kestrel-bay-field-guide.docx"
DECK = "orbit-sensor-launch-deck.pptx"
RUNBOOK = "incident-runbook.md"


def _font(size: int) -> Any:
    try:
        return ImageFont.load_default(size=size)
    except TypeError:  # older Pillow
        return ImageFont.load_default()


def bar_chart_png(title: str, labels: list[str], values: list[float], unit: str = "") -> bytes:
    w, h = 900, 560
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((40, 24), title, fill="black", font=_font(30))
    left, right, top, bottom = 90, w - 40, 100, h - 80
    d.line((left, bottom, right, bottom), fill="black", width=2)
    d.line((left, top, left, bottom), fill="black", width=2)
    vmax = max(values) * 1.15
    slot = (right - left) / len(values)
    for i, (label, value) in enumerate(zip(labels, values, strict=True)):
        x0 = left + i * slot + slot * 0.2
        x1 = left + (i + 1) * slot - slot * 0.2
        y0 = bottom - (bottom - top) * value / vmax
        d.rectangle((x0, y0, x1, bottom), fill=(255, 86, 35))
        d.text((x0, y0 - 34), f"{value:g}{unit}", fill="black", font=_font(24))
        d.text((x0, bottom + 14), label, fill="black", font=_font(24))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _report_pdf(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    doc.set_metadata({"title": "Halcyon Grid Co. Annual Report 2025"})
    body = pymupdf.Font("helv")

    def write(page, y: float, text: str, size: float = 11, bold: bool = False) -> float:
        rect = pymupdf.Rect(60, y, 552, y + 400)
        page.insert_textbox(rect, text, fontsize=size, fontname="hebo" if bold else "helv")
        lines = max(1, int(body.text_length(text, fontsize=size) / 480) + text.count("\n") + 1)
        return y + lines * size * 1.45 + 10

    p1 = doc.new_page()
    y = write(p1, 60, "Halcyon Grid Co. Annual Report 2025", 22, True)
    y = write(p1, y + 6, "Letter from the Chief Executive", 14, True)
    y = write(
        p1,
        y,
        (
            "2025 was the year Halcyon Grid Co. became a storage company as much as a generation company. "
            "Total revenue reached 412 million euros, up 18 percent on 2024, driven by the Saltmarsh battery "
            "park in the Netherlands, which entered commercial operation in March. Chief Executive Ines Varga "
            "said the company would prioritise grid-scale storage over new gas capacity through 2028."
        ),
    )
    y = write(
        p1,
        y,
        (
            "Net carbon intensity fell to 96 grams of CO2 per kilowatt-hour, compared with 141 grams in 2024. "
            "The board approved a dividend of 0.42 euros per share. Employee headcount grew to 1,870 people, "
            "and the lost-time injury rate improved to 0.31 per 200,000 hours worked."
        ),
    )
    write(p1, y, "Outlook", 14, True)
    write(
        p1,
        y + 30,
        (
            "For 2026 the company guides revenue between 450 and 480 million euros. The largest risk named by "
            "management is grid connection delays for the Dunlin offshore wind project, now expected in 2027."
        ),
    )

    p2 = doc.new_page()
    y = write(p2, 60, "Segment results", 16, True)
    y = write(p2, y, "Table 2: Revenue by segment, million euros. Storage was the fastest-growing segment.")
    rows = [
        ["Segment", "2024", "2025", "Change"],
        ["Wind", "188", "201", "+7%"],
        ["Solar", "97", "104", "+7%"],
        ["Storage", "35", "79", "+126%"],
        ["Retail", "29", "28", "-3%"],
        ["Total", "349", "412", "+18%"],
    ]
    x0, col, row_h = 60, 120, 24
    top = y + 6
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            rect = pymupdf.Rect(x0 + c * col, top + r * row_h, x0 + (c + 1) * col, top + (r + 1) * row_h)
            p2.draw_rect(rect, color=(0, 0, 0), width=0.8)
            p2.insert_textbox(
                pymupdf.Rect(rect.x0 + 6, rect.y0 + 5, rect.x1, rect.y1),
                cell,
                fontsize=10,
                fontname="hebo" if r == 0 else "helv",
            )
    y = top + len(rows) * row_h + 24
    write(
        p2,
        y,
        (
            "Storage revenue more than doubled because Saltmarsh sold frequency-response services to the "
            "Dutch transmission operator. Retail revenue declined slightly after the company exited two "
            "low-margin municipal supply contracts."
        ),
    )

    p3 = doc.new_page()
    y = write(p3, 60, "Generation mix", 16, True)
    chart = bar_chart_png(
        "Installed capacity by technology (MW), 2025",
        ["Wind", "Solar", "Storage", "Gas"],
        [1240, 610, 455, 300],
    )
    rect = pymupdf.Rect(60, y + 10, 552, y + 10 + 492 * 560 / 900)
    p3.insert_image(rect, stream=chart)
    write(p3, rect.y1 + 8, "Figure 3: Installed capacity by technology in megawatts at year end 2025.", 10)
    write(
        p3,
        rect.y1 + 40,
        (
            "Gas capacity will be reduced to 150 megawatts by 2027 as the Ridgeback plant is mothballed. "
            "No new gas investments are planned."
        ),
    )
    doc.save(path)
    doc.close()


def _field_guide_docx(path: Path) -> None:
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.core_properties.title = "Kestrel Bay Coastal Field Guide"
    d.add_heading("Kestrel Bay Coastal Field Guide", 0)
    d.add_heading("About the reserve", 1)
    d.add_paragraph(
        "Kestrel Bay Nature Reserve covers 14 square kilometres of saltmarsh, dunes and mudflats. "
        "It is managed by the Kestrel Bay Trust, which counts 212 volunteer wardens."
    )
    d.add_paragraph(
        "The best time for migrating waders is late September, two hours either side of high tide. "
        "Visitors must stay on marked boardwalks between April and July to protect nesting terns."
    )
    d.add_heading("Species counts", 1)
    d.add_paragraph("The 2025 winter census recorded the following peak counts.")
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(["Species", "Peak count", "Trend vs 2024"]):
        t.rows[0].cells[i].text = h
    for row in [
        ("Dunlin", "4,300", "up"),
        ("Bar-tailed godwit", "1,150", "stable"),
        ("Brent goose", "2,750", "down"),
        ("Avocet", "640", "up"),
    ]:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    d.add_heading("Seasonal visitors", 1)
    d.add_picture(
        io.BytesIO(
            bar_chart_png(
                "Visitors per season (thousands)", ["Spring", "Summer", "Autumn", "Winter"], [38, 61, 27, 12]
            )
        ),
        width=Inches(5.5),
    )
    d.add_paragraph("Figure 2: Visitor numbers peak in summer at 61 thousand.")
    d.add_heading("Safety", 1)
    d.add_paragraph(
        "Tides rise faster than walking pace across the eastern mudflats. Check the tide table at "
        "the visitor centre and never cross the Heron Channel within three hours of high water."
    )
    d.save(str(path))


def _deck_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.core_properties.title = "Orbit S2 Sensor Launch"
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Orbit S2 Sensor Launch"
    s.placeholders[1].text = "Go-to-market plan, Q1 2026"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Pricing and availability"
    s.placeholders[1].text = (
        "Orbit S2 ships on 12 February 2026.\nList price is 249 dollars, "
        "or 219 dollars with a two-year monitoring plan.\nBattery life is 26 months."
    )
    s.notes_slide.notes_text_frame.text = "Stress that the S2 replaces the S1 at the same price point."
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Pre-orders by region"
    data = CategoryChartData()
    data.categories = ["North America", "Europe", "APAC"]
    data.add_series("Pre-orders (units)", (8200, 5100, 2600))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(8), Inches(4.5), data)
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Launch risks"
    rows, cols = 4, 3
    table = s.shapes.add_table(rows, cols, Inches(0.6), Inches(1.6), Inches(8.8), Inches(2.4)).table
    for r, row in enumerate(
        [
            ("Risk", "Owner", "Mitigation"),
            ("Chip supply", "Dana Okafor", "Second supplier"),
            ("Certification delay", "Luis Brandt", "Pre-test in December"),
            ("App store review", "Mei Tanaka", "Submit two weeks early"),
        ]
    ):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
            table.cell(r, c).text_frame.paragraphs[0].font.size = Pt(14)
    prs.save(str(path))


RUNBOOK_TEXT = """# Incident Runbook: Payments API

## Severity levels

| Severity | Definition | Response time |
|---|---|---|
| SEV1 | Payments failing for more than 5% of requests | 5 minutes |
| SEV2 | Elevated latency above 800 ms at p95 | 15 minutes |
| SEV3 | Single merchant affected | 4 hours |

## First steps

The on-call engineer acknowledges the page in PagerDuty, opens a channel named inc-payments-<date>, and
assigns an incident commander. For SEV1 incidents the commander must notify the status page within 10 minutes.

## Rollback

Deployments are rolled back with the `deployctl rollback payments --to previous` command. A rollback is
preferred over a forward fix whenever the faulty change shipped less than two hours ago.
"""


def write_samples(target: Path) -> list[Path]:
    target.mkdir(parents=True, exist_ok=True)
    paths = [target / REPORT, target / FIELD_GUIDE, target / DECK, target / RUNBOOK]
    _report_pdf(paths[0])
    _field_guide_docx(paths[1])
    _deck_pptx(paths[2])
    paths[3].write_text(RUNBOOK_TEXT)
    return paths


GOLDEN: list[dict[str, Any]] = [
    {
        "question": "What was Halcyon Grid's total revenue in 2025?",
        "reference_answer": "412 million euros, up 18%.",
        "expected_document": REPORT,
        "expected_page": 1,
        "tag": "text",
    },
    {
        "question": "How much revenue did the storage segment make in 2025 and how did it change?",
        "reference_answer": "79 million euros, up 126% from 35 million.",
        "expected_document": REPORT,
        "expected_page": 2,
        "tag": "table",
    },
    {
        "question": "How many megawatts of wind capacity were installed at the end of 2025?",
        "reference_answer": "1,240 MW of wind.",
        "expected_document": REPORT,
        "expected_page": 3,
        "tag": "figure",
    },
    {
        "question": "What revenue range does Halcyon Grid guide for 2026?",
        "reference_answer": "Between 450 and 480 million euros.",
        "expected_document": REPORT,
        "expected_page": 1,
        "tag": "text",
    },
    {
        "question": "What was the peak dunlin count in the winter census?",
        "reference_answer": "4,300, trending up.",
        "expected_document": FIELD_GUIDE,
        "expected_page": None,
        "tag": "table",
    },
    {
        "question": "When is the best time to see migrating waders at Kestrel Bay?",
        "reference_answer": "Late September, two hours either side of high tide.",
        "expected_document": FIELD_GUIDE,
        "expected_page": None,
        "tag": "text",
    },
    {
        "question": "What is the list price of the Orbit S2 and when does it ship?",
        "reference_answer": "249 dollars (219 with a two-year plan); ships 12 February 2026.",
        "expected_document": DECK,
        "expected_page": 2,
        "tag": "text",
    },
    {
        "question": "Which region had the most Orbit S2 pre-orders?",
        "reference_answer": "North America with 8,200 units.",
        "expected_document": DECK,
        "expected_page": 3,
        "tag": "table",
    },
    {
        "question": "Who owns the certification delay risk for the launch?",
        "reference_answer": "Luis Brandt; mitigation is pre-testing in December.",
        "expected_document": DECK,
        "expected_page": 4,
        "tag": "table",
    },
    {
        "question": "What is the response time for a SEV2 payments incident?",
        "reference_answer": "15 minutes.",
        "expected_document": RUNBOOK,
        "expected_page": None,
        "tag": "table",
    },
    {
        "question": "How do you roll back the payments service?",
        "reference_answer": "Run deployctl rollback payments --to previous.",
        "expected_document": RUNBOOK,
        "expected_page": None,
        "tag": "text",
    },
]
